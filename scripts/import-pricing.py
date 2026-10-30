"""Copies the van placeholder prices out of the master pricing workbook.

    python3 scripts/import-pricing.py "<Dropbox>/Papago Vans/van build pricing/papago-master-pricing.xlsx"

The workbook (owner, 2026-10-30) is the single source of truth for every
advertised price. This writes lib/pricing-data.json, which the builder reads.
Retail numbers only: costs and markup never leave the workbook, because this
repo is public. It carries the van prices (Chassis tab) and the base options
retail for a short and a long van (Base Options tab); upgrades follow.

Retail is worked out here from the workbook's inputs (qty, unit cost, install
hours, markup, Settings) with the same arithmetic as its formulas, because a
workbook saved by a script holds no calculated values until Excel saves it.

Needs openpyxl (pip3 install openpyxl).
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

# Workbook chassis code -> the builder's van id (the ?b= links use these).
VAN_IDS = {
    "S144": "sprinter-144",
    "S170": "sprinter-170",
    "S170X": "sprinter-170x",
    "T148": "transit-148",
    "T148X": "transit-148x",
    "P136": "promaster-136",
    "P159": "promaster-159",
    "P159X": "promaster-159x",
}

wb = load_workbook(sys.argv[1], data_only=True)
ws = wb["Chassis"]
head = {c.value: c.column for c in ws[4] if c.value}
code_col, price_col = head["Code"], head["Placeholder Price"]

vans = {}
for row in ws.iter_rows(min_row=5):
    code = row[code_col - 1].value
    if code in ("Van Options", "Code"):
        break  # the Van Options table starts here
    if not code:
        continue
    price = row[price_col - 1].value
    if code not in VAN_IDS:
        sys.exit(f"{code}: no builder id in VAN_IDS; add one")
    if not isinstance(price, (int, float)):
        sys.exit(f"{code}: Placeholder Price is not a number ({price!r})")
    vans[VAN_IDS[code]] = {"code": code, "price": round(price)}

missing = set(VAN_IDS.values()) - set(vans)
if missing:
    sys.exit(f"workbook has no row for {', '.join(sorted(missing))}")

# Base options retail: qty x (unit cost + install hrs x labor) x (1 + markup),
# summed over every row that fits All, plus the rows for that size.
st = wb["Settings"]
markup, labor = st["B4"].value or 0, st["B5"].value or 0
bo = wb["Base Options"]
cols = {c.value: c.column - 1 for c in bo[4] if c.value}
base = {"short": 0.0, "long": 0.0}
for row in bo.iter_rows(min_row=5, values_only=True):
    if not row[cols["Code"]]:
        continue
    cost = (row[cols["Qty"]] or 0) * ((row[cols["Unit Cost"]] or 0) + (row[cols["Install Hrs"]] or 0) * labor)
    m = row[cols["Markup %"]]
    retail = cost * (1 + (markup if m is None else m))
    fits = row[cols["Fits"]]
    for size in base:
        if fits in ("All", size.capitalize()):
            base[size] += retail
base = {k: round(v) for k, v in base.items()}
if not base["short"]:
    sys.exit("Base Options retail came to $0: fill in the costs first")

out = Path(__file__).resolve().parent.parent / "lib" / "pricing-data.json"
out.write_text(json.dumps({"source": Path(sys.argv[1]).name, "vans": vans, "base": base}, indent=2) + "\n")
for vid, v in vans.items():
    print(f"{v['code']:6} {vid:15} ${v['price']:,}")
print(f"base options retail: short ${base['short']:,}, long ${base['long']:,}")
print(f"wrote {out}")
